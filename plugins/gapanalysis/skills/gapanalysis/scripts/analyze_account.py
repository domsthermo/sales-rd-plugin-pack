#!/usr/bin/env python3
"""Deterministic customer purchase-history and gap analysis.

The script reads a source workbook without modifying it, resolves a narrow
account scope, calculates matched-period movement, identifies mechanical lapse
candidates, maps observed purchases to a small workflow catalog, and writes
plain CSV/JSON/Markdown outputs.  It intentionally does not make pricing,
probability, contact, outreach, or CRM decisions.
"""

from __future__ import annotations

import argparse
import calendar
import csv
from dataclasses import dataclass
from datetime import date, datetime
from difflib import SequenceMatcher
import json
import math
from pathlib import Path
import re
import statistics
import sys
from typing import Any, Iterable


VERSION = "1.0.0"

FIELD_ALIASES: dict[str, list[str]] = {
    "invoice_date": ["invoice date", "invoice dt", "billing date", "date"],
    "order_date": ["order date", "sales order date"],
    "account_name": ["ship to", "account name", "customer name", "account", "customer"],
    "account_id": ["nsgn", "account id", "account number", "account no", "customer id", "customer number"],
    "site_id": ["sgn", "site sgn", "site id", "site number"],
    "sku": ["sku", "catalog number", "catalog no", "catalog #", "product number", "item number", "material number"],
    "product_description": ["product description", "sku description", "item description", "material description", "description"],
    "product_family": ["pricing product line", "product line group", "dsr product group na", "product family", "product line", "product group"],
    "quantity": ["ship quantity", "invoice quantity", "quantity", "qty", "ordered quantity"],
    "unit_price": ["net unit price", "unit price", "price each"],
    "net_sales": ["net sales", "net revenue", "invoice net", "sales net"],
    "extended_sales": ["ext sales", "extended sales", "extended amount", "sales amount", "revenue"],
    "freight": ["freight", "shipping", "shipping charge"],
    "currency": ["currency", "currency code", "curr"],
    "invoice_number": ["sales invoice #", "sales invoice number", "sales invoice no", "invoice number", "invoice #"],
}

REQUIRED_FIELD_GROUPS = (
    ("invoice_date",),
    ("net_sales", "extended_sales"),
    ("sku", "product_description", "product_family"),
)

GENERIC_ACCOUNT_TOKENS = {
    "bio", "biotech", "biotherapeutics", "biosciences", "company", "inc",
    "lab", "labs", "life", "pharma", "pharmaceuticals", "sciences", "therapeutics",
}

LEGAL_SUFFIXES = {"co", "company", "corp", "corporation", "inc", "incorporated", "llc", "lp", "ltd", "limited", "plc"}

LAPSE_EXCLUSIONS = (
    "instrument", "cytometer", "microscope", "analyzer", "reader", "sequencer",
    "thermal cycler", "quantstudio", "attune", "evos", "service", "maintenance",
    "freight", "shipping", "delivery", "installation", "project", "software", "license",
)

OPERATIONAL_TERMS = (
    "freight", "shipping charge", "handling charge", "delivery charge",
    "hazardous charge", "dry ice charge", "sales tax",
)


@dataclass
class SourceProfile:
    sheet: str
    header_row: int
    mapping: dict[str, int]
    headers: list[str]
    row_count: int
    format: str


def norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\u2019", "'").casefold().strip()
    return re.sub(r"\s+", " ", text)


def token_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", norm(value)).strip()


def clean_identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\.0$", "", str(value).strip())


def parse_number(value: Any) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        result = float(value)
        return result if math.isfinite(result) else None
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"[$,£€%()]", "", text).strip()
    if not text:
        return None
    try:
        result = float(text)
    except ValueError:
        return None
    if negative:
        result = -result
    return result if math.isfinite(result) else None


DATE_FORMATS = (
    "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%d-%b-%Y", "%d-%b-%y",
)


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        compact = str(int(value)) if float(value).is_integer() else ""
        if len(compact) == 8 and compact.isdigit():
            try:
                return datetime.strptime(compact, "%Y%m%d").date()
            except ValueError:
                pass
        try:
            from openpyxl.utils.datetime import from_excel
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except Exception:
            return None
    text = str(value).strip()
    if re.fullmatch(r"\d{8}", text):
        try:
            return datetime.strptime(text, "%Y%m%d").date()
        except ValueError:
            pass
    for format_string in DATE_FORMATS:
        try:
            return datetime.strptime(text, format_string).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def resolve_mapping(headers: list[Any]) -> dict[str, int]:
    normalized = [token_text(value) for value in headers]
    mapping: dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            matches = [index for index, header in enumerate(normalized) if header == token_text(alias)]
            if len(matches) == 1:
                mapping[field] = matches[0]
                break
    return mapping


def header_score(mapping: dict[str, int]) -> int:
    score = len(mapping) * 5
    score += 25 if "invoice_date" in mapping else 0
    score += 25 if "net_sales" in mapping or "extended_sales" in mapping else 0
    score += 15 if "sku" in mapping else 0
    score += 12 if "product_family" in mapping else 0
    score += 15 if "site_id" in mapping else 0
    score += 8 if "account_name" in mapping else 0
    return score


def validate_mapping(mapping: dict[str, int]) -> None:
    missing = [" or ".join(group) for group in REQUIRED_FIELD_GROUPS if not any(field in mapping for field in group)]
    if missing:
        raise ValueError("Missing required column group(s): " + ", ".join(missing))
    if not any(field in mapping for field in ("site_id", "account_id", "account_name")):
        raise ValueError("No account field found. Expected SGN, NSGN, Ship To, Account, or Customer.")


def csv_rows(path: Path) -> tuple[SourceProfile, list[list[Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        rows = list(csv.reader(source))
    if not rows:
        raise ValueError("The CSV file is empty.")
    candidates = []
    for index, row in enumerate(rows[:30], start=1):
        mapping = resolve_mapping(row)
        candidates.append((header_score(mapping), -index, index, mapping, row))
    score, _, header_row, mapping, headers = max(candidates)
    if score < 45:
        raise ValueError("No credible purchase-history header row was found in the CSV file.")
    validate_mapping(mapping)
    profile = SourceProfile(path.stem, header_row, mapping, [str(value) for value in headers], len(rows), "csv")
    return profile, rows[header_row:]


def workbook_rows(path: Path, requested_sheet: str | None) -> tuple[SourceProfile, list[list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("openpyxl is required to read Excel files in this environment.") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        profiles = []
        for sheet_name in workbook.sheetnames:
            if requested_sheet and sheet_name != requested_sheet:
                continue
            worksheet = workbook[sheet_name]
            for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=min(30, worksheet.max_row), values_only=True), start=1):
                headers = list(row)
                mapping = resolve_mapping(headers)
                profiles.append((header_score(mapping), -row_number, sheet_name, row_number, mapping, headers, worksheet.max_row))
        if requested_sheet and requested_sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {requested_sheet}")
        if not profiles:
            raise ValueError("The workbook does not contain a readable worksheet.")
        best_score = max(item[0] for item in profiles)
        best = [item for item in profiles if item[0] == best_score]
        best_sheets = sorted({item[2] for item in best})
        if best_score < 45:
            raise ValueError("No worksheet contains a credible purchase-history header row.")
        if len(best_sheets) > 1 and not requested_sheet:
            raise ValueError("Multiple worksheets are equally plausible. Specify --sheet: " + ", ".join(best_sheets))
        _, _, sheet_name, header_row, mapping, headers, max_row = sorted(best, key=lambda item: (item[2], -item[1]))[0]
        validate_mapping(mapping)
        worksheet = workbook[sheet_name]
        data = [list(row) for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True)]
        profile = SourceProfile(sheet_name, header_row, mapping, [str(value or "") for value in headers], max_row, "xlsx")
        return profile, data
    finally:
        workbook.close()


def read_source(path: Path, requested_sheet: str | None) -> tuple[SourceProfile, list[list[Any]]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        if requested_sheet:
            raise ValueError("--sheet cannot be used with a CSV file.")
        return csv_rows(path)
    if suffix in {".xlsx", ".xlsm"}:
        return workbook_rows(path, requested_sheet)
    raise ValueError("Supported input formats are .xlsx, .xlsm, and .csv.")


def cell(row: list[Any], mapping: dict[str, int], field: str) -> Any:
    index = mapping.get(field)
    return row[index] if index is not None and index < len(row) else None


def build_records(profile: SourceProfile, rows: list[list[Any]]) -> tuple[list[dict[str, Any]], dict[str, int], str]:
    measure_field = "net_sales" if "net_sales" in profile.mapping else "extended_sales"
    records: list[dict[str, Any]] = []
    counters = {"source_rows": len(rows), "blank_rows": 0, "missing_sales": 0, "missing_date": 0}
    for offset, row in enumerate(rows, start=profile.header_row + 1):
        selected_values = [cell(row, profile.mapping, field) for field in profile.mapping]
        if not any(value not in (None, "") for value in selected_values):
            counters["blank_rows"] += 1
            continue
        sales = parse_number(cell(row, profile.mapping, measure_field))
        invoice_date = parse_date(cell(row, profile.mapping, "invoice_date"))
        if sales is None:
            counters["missing_sales"] += 1
        if invoice_date is None:
            counters["missing_date"] += 1
        sku = clean_identifier(cell(row, profile.mapping, "sku"))
        description = clean_identifier(cell(row, profile.mapping, "product_description"))
        if not description and " - " in sku:
            sku, description = sku.split(" - ", 1)
        records.append({
            "source_row": offset,
            "invoice_date": invoice_date,
            "account_name": clean_identifier(cell(row, profile.mapping, "account_name")),
            "account_id": clean_identifier(cell(row, profile.mapping, "account_id")),
            "site_id": clean_identifier(cell(row, profile.mapping, "site_id")),
            "sku": sku,
            "product_description": description,
            "product_family": clean_identifier(cell(row, profile.mapping, "product_family")) or "Unmapped",
            "quantity": parse_number(cell(row, profile.mapping, "quantity")),
            "sales": sales,
            "currency": clean_identifier(cell(row, profile.mapping, "currency")),
            "invoice_number": clean_identifier(cell(row, profile.mapping, "invoice_number")),
        })
    return records, counters, measure_field


def parse_sgn_label(value: str) -> tuple[str, str]:
    match = re.match(r"^\s*(\d+)\s*[-–—]\s*(.*?)\s*$", value)
    if match:
        return match.group(1), match.group(2)
    return (value if value.isdigit() else "", "" if value.isdigit() else value)


def account_key(record: dict[str, Any]) -> str:
    return record.get("site_id") or record.get("account_id") or record.get("account_name") or ""


def account_identities(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    for record in records:
        key = account_key(record)
        if not key:
            continue
        item = grouped.setdefault(key, {"key": key, "row_count": 0, "sales": 0.0, "names": {}})
        item["row_count"] += 1
        if record.get("sales") is not None:
            item["sales"] += float(record["sales"])
        name = record.get("account_name") or ""
        if name:
            item["names"][name] = item["names"].get(name, 0) + 1
    identities = []
    for key, item in grouped.items():
        numeric, embedded_name = parse_sgn_label(key)
        name = embedded_name
        if not name and item["names"]:
            name = sorted(item["names"].items(), key=lambda pair: (-pair[1], pair[0].casefold()))[0][0]
        label = key if embedded_name else f"{key} - {name}".strip(" -")
        identities.append({
            "key": key,
            "numeric_id": numeric or (key if key.isdigit() else ""),
            "name": name or key,
            "label": label,
            "row_count": item["row_count"],
            "sales": round(item["sales"], 2),
        })
    return sorted(identities, key=lambda item: (-item["row_count"], item["label"].casefold()))


def company_tokens(value: str) -> list[str]:
    _, name = parse_sgn_label(value)
    tokens = re.findall(r"[a-z0-9]+", norm(name or value))
    while tokens and tokens[-1] in LEGAL_SUFFIXES:
        tokens.pop()
    return tokens


def match_score(query_tokens: list[str], candidate_tokens: list[str]) -> float:
    if not query_tokens or not candidate_tokens:
        return 0.0
    query = " ".join(query_tokens)
    candidate = " ".join(candidate_tokens)
    if query in candidate:
        return 1.0
    span = min(len(query_tokens), len(candidate_tokens))
    scores = [SequenceMatcher(None, query, candidate).ratio()]
    scores.extend(
        SequenceMatcher(None, query, " ".join(candidate_tokens[index:index + span])).ratio()
        for index in range(len(candidate_tokens) - span + 1)
    )
    if len(query_tokens) == 1:
        scores.extend(SequenceMatcher(None, query, token).ratio() for token in candidate_tokens)
    return max(scores)


def common_name(matches: list[dict[str, Any]], fallback: str) -> str:
    token_lists = [company_tokens(item["name"]) for item in matches]
    shared: list[str] = []
    for group in zip(*token_lists):
        if len(set(group)) != 1:
            break
        shared.append(group[0])
    if shared:
        return " ".join(shared).upper()
    if len(matches) == 1:
        return matches[0]["name"]
    return fallback.strip()


def resolve_accounts(query: str | None, identities: list[dict[str, Any]]) -> dict[str, Any]:
    if not identities:
        raise ValueError("No account identifiers were found in the purchase rows.")
    if not query:
        if len(identities) == 1:
            return {"method": "only_account", "display_name": identities[0]["name"], "matches": identities}
        suggestions = " | ".join(item["label"] for item in identities[:15])
        raise ValueError("The file contains multiple accounts. Supply --customer. Largest account scopes: " + suggestions)
    requested = norm(query)
    exact = [item for item in identities if requested in {norm(item["key"]), norm(item["label"]), norm(item["numeric_id"])}]
    if len(exact) == 1:
        return {"method": "exact_account", "display_name": exact[0]["name"], "matches": exact}
    query_tokens = company_tokens(query)
    if not query_tokens:
        raise ValueError("The customer query does not contain a usable name or account identifier.")
    if len(query_tokens) == 1 and query_tokens[0] in GENERIC_ACCOUNT_TOKENS:
        raise ValueError(f"Customer query {query!r} is too broad for safe account matching.")
    scored = []
    for item in identities:
        score = match_score(query_tokens, company_tokens(f"{item['label']} {item['name']}"))
        scored.append({**item, "match_score": round(score, 4)})
    scored.sort(key=lambda item: (-item["match_score"], -item["row_count"], item["label"].casefold()))
    best = scored[0]["match_score"]
    compact = "".join(query_tokens)
    threshold = 0.70 if len(compact) >= 6 else 0.82
    if best < threshold:
        suggestions = " | ".join(item["label"] for item in scored[:8])
        raise ValueError(f"No account closely matches {query!r}. Nearest accounts: {suggestions}")
    selected = [item for item in scored if item["match_score"] >= threshold]
    if len(selected) > 20:
        raise ValueError(f"Customer query {query!r} matched {len(selected)} accounts. Use a more specific name or exact account/SGN.")
    return {"method": "fuzzy_company_rollup", "display_name": common_name(selected, query), "matches": selected}


def prior_cutoff(latest_cutoff: date) -> date:
    year = latest_cutoff.year - 1
    day = min(latest_cutoff.day, calendar.monthrange(year, latest_cutoff.month)[1])
    return date(year, latest_cutoff.month, day)


def period_flags(value: date | None, latest_cutoff: date) -> tuple[bool, bool]:
    if value is None:
        return False, False
    previous_cutoff = prior_cutoff(latest_cutoff)
    latest = value.year == latest_cutoff.year and value <= latest_cutoff
    prior = value.year == previous_cutoff.year and value <= previous_cutoff
    return latest, prior


def pct_change(current: float, prior: float) -> float | None:
    return round((current - prior) / abs(prior), 6) if prior else None


def movement(current: float, prior: float) -> str:
    epsilon = 0.005
    if current > epsilon and prior <= epsilon:
        return "gained"
    if prior > epsilon and current <= epsilon:
        return "lost"
    difference = current - prior
    if abs(difference) <= epsilon:
        return "stable"
    return "growing" if difference > 0 else "declining"


def is_operational(record: dict[str, Any]) -> bool:
    family = norm(record.get("product_family"))
    text = norm(f"{record.get('product_family', '')} {record.get('product_description', '')} {record.get('sku', '')}")
    return family.startswith("fr1 ") or family.startswith("fr4 ") or any(term in text for term in OPERATIONAL_TERMS)


def aggregate_trends(records: list[dict[str, Any]], latest_cutoff: date, key_fields: tuple[str, ...]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    for record in records:
        if record.get("sales") is None or record.get("invoice_date") is None:
            continue
        key = tuple(str(record.get(field) or "") for field in key_fields)
        if not any(key):
            continue
        item = grouped.setdefault(key, {
            **{field: key[index] for index, field in enumerate(key_fields)},
            "latest_ytd_sales": 0.0,
            "prior_ytd_sales": 0.0,
            "latest_ytd_quantity": 0.0,
            "prior_ytd_quantity": 0.0,
            "latest_ytd_rows": 0,
            "prior_ytd_rows": 0,
        })
        is_latest, is_prior = period_flags(record["invoice_date"], latest_cutoff)
        quantity = float(record.get("quantity") or 0.0)
        if is_latest:
            item["latest_ytd_sales"] += float(record["sales"])
            item["latest_ytd_quantity"] += quantity
            item["latest_ytd_rows"] += 1
        if is_prior:
            item["prior_ytd_sales"] += float(record["sales"])
            item["prior_ytd_quantity"] += quantity
            item["prior_ytd_rows"] += 1
    output = []
    for item in grouped.values():
        current = round(item["latest_ytd_sales"], 2)
        previous = round(item["prior_ytd_sales"], 2)
        item["latest_ytd_sales"] = current
        item["prior_ytd_sales"] = previous
        item["latest_ytd_quantity"] = round(item["latest_ytd_quantity"], 3)
        item["prior_ytd_quantity"] = round(item["prior_ytd_quantity"], 3)
        item["change"] = round(current - previous, 2)
        item["change_percent"] = pct_change(current, previous)
        item["movement"] = movement(current, previous)
        output.append(item)
    return sorted(output, key=lambda item: (-abs(item["change"]), -max(abs(item["latest_ytd_sales"]), abs(item["prior_ytd_sales"])), tuple(item.get(field, "") for field in key_fields)))


def lapsed_candidates(records: list[dict[str, Any]], latest_cutoff: date) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str], dict[date, float]] = {}
    for record in records:
        sku = record.get("sku") or ""
        purchase_date = record.get("invoice_date")
        sales = record.get("sales")
        if not sku or purchase_date is None or sales is None or float(sales) <= 0 or purchase_date > latest_cutoff:
            continue
        key = (sku, record.get("product_description") or "", record.get("product_family") or "Unmapped")
        grouped.setdefault(key, {})[purchase_date] = grouped.setdefault(key, {}).get(purchase_date, 0.0) + float(sales)
    output = []
    for (sku, description, family), by_date in grouped.items():
        dates = sorted(by_date)
        if len(dates) < 3:
            continue
        latest_positive_in_period = any(period_flags(value, latest_cutoff)[0] for value in dates)
        if latest_positive_in_period:
            continue
        text = norm(f"{description} {family}")
        if any(term in text for term in LAPSE_EXCLUSIONS):
            continue
        intervals = [(dates[index] - dates[index - 1]).days for index in range(1, len(dates))]
        if len(intervals) < 2:
            continue
        median_cadence = statistics.median(intervals)
        days_since = (latest_cutoff - dates[-1]).days
        threshold = max(90.0, 1.75 * median_cadence)
        if days_since <= threshold:
            continue
        output.append({
            "sku": sku,
            "product_description": description,
            "product_family": family,
            "distinct_purchase_dates": len(dates),
            "last_positive_purchase": dates[-1].isoformat(),
            "median_cadence_days": round(median_cadence, 1),
            "days_since_last_purchase": days_since,
            "historical_positive_sales": round(sum(by_date.values()), 2),
            "status": "Win-back hypothesis; confirm workflow continuity and account/site migration.",
        })
    return sorted(output, key=lambda item: (-item["historical_positive_sales"], -item["days_since_last_purchase"], item["sku"]))


def load_catalog(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict) or not isinstance(value.get("categories"), dict) or not isinstance(value.get("workflows"), list):
        raise ValueError(f"Invalid workflow catalog: {path}")
    return value


def load_category_map(path: Path | None, valid_categories: set[str]) -> list[tuple[str, str]]:
    if path is None:
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if not reader.fieldnames or not {"match", "category"}.issubset({norm(value) for value in reader.fieldnames}):
            raise ValueError("Category map needs columns named match and category.")
        normalized_names = {norm(name): name for name in reader.fieldnames}
        output = []
        for row in reader:
            match = token_text(row.get(normalized_names["match"]))
            category = str(row.get(normalized_names["category"]) or "").strip()
            if not match or not category:
                continue
            if category not in valid_categories:
                raise ValueError(f"Unknown category in category map: {category}")
            output.append((match, category))
    return sorted(output, key=lambda item: -len(item[0]))


def classify_category(record: dict[str, Any], catalog: dict[str, Any], overrides: list[tuple[str, str]]) -> str | None:
    text = token_text(f"{record.get('product_family', '')} {record.get('product_description', '')} {record.get('sku', '')}")
    if not text:
        return None
    for match, category in overrides:
        if match in text:
            return category
    candidates = []
    for category, keywords in catalog["categories"].items():
        matched = [token_text(keyword) for keyword in keywords if token_text(keyword) and token_text(keyword) in text]
        if matched:
            candidates.append((max(len(value) for value in matched), len(matched), category))
    return sorted(candidates, reverse=True)[0][2] if candidates else None


def category_evidence(records: list[dict[str, Any]], latest_cutoff: date, catalog: dict[str, Any], overrides: list[tuple[str, str]]) -> tuple[set[str], set[str], dict[str, float], dict[str, float], dict[str, float]]:
    current: dict[str, float] = {}
    historical: dict[str, float] = {}
    unmapped: dict[str, float] = {}
    for record in records:
        sales = record.get("sales")
        invoice_date = record.get("invoice_date")
        if sales is None:
            continue
        amount = float(sales)
        category = classify_category(record, catalog, overrides)
        if category:
            historical[category] = historical.get(category, 0.0) + amount
            if invoice_date and period_flags(invoice_date, latest_cutoff)[0]:
                current[category] = current.get(category, 0.0) + amount
        else:
            family = record.get("product_family") or "Unmapped"
            unmapped[family] = unmapped.get(family, 0.0) + amount
    historical_positive = {key for key, value in historical.items() if value > 0}
    current_positive = {key for key, value in current.items() if value > 0}
    return current_positive, historical_positive, current, historical, unmapped


def load_research(path: Path | None) -> dict[str, Any]:
    if path is None:
        return {"company_summary": "", "workflows": [], "sources": []}
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("Research input must be a JSON object.")
    if not isinstance(value.get("workflows", []), list) or not isinstance(value.get("sources", []), list):
        raise ValueError("Research input workflows and sources must be arrays.")
    return value


def build_gaps(catalog: dict[str, Any], current_observed: set[str], historical_observed: set[str], research: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    workflows_by_id = {item["id"]: item for item in catalog["workflows"]}
    research_by_id = {}
    warnings = []
    for item in research.get("workflows", []):
        workflow_id = str(item.get("id") or "").strip()
        if not workflow_id:
            continue
        if workflow_id not in workflows_by_id:
            warnings.append(f"Research workflow id ignored because it is not in the catalog: {workflow_id}")
            continue
        research_by_id[workflow_id] = item
    gaps = []
    for workflow in catalog["workflows"]:
        workflow_id = workflow["id"]
        research_item = research_by_id.get(workflow_id)
        current_anchors = [category for category in workflow.get("anchors", []) if category in current_observed]
        historical_anchors = [category for category in workflow.get("anchors", []) if category in historical_observed]
        anchors = current_anchors or historical_anchors
        if not anchors and not research_item:
            continue
        if anchors and research_item:
            basis = "Observed purchase anchor + public research"
        elif anchors:
            basis = "Observed purchase adjacency"
        else:
            basis = "Public workflow research"
        missing = [category for category in workflow.get("categories", []) if category not in historical_observed]
        for category in missing[:5]:
            gaps.append({
                "workflow": workflow["name"],
                "evidence_basis": basis,
                "observed_anchor_categories": "; ".join(value.replace("-", " ") for value in anchors),
                "missing_category": category.replace("-", " "),
                "gap_statement": "Not observed in the supplied purchase history for the matched account scope and period.",
                "research_evidence": str((research_item or {}).get("evidence") or ""),
                "source_title": str((research_item or {}).get("source_title") or ""),
                "source_url": str((research_item or {}).get("source_url") or ""),
                "validation_question": f"Is {category.replace('-', ' ')} currently used in the {workflow['name'].casefold()} workflow, and if so, where is it sourced?",
            })
    basis_order = {
        "Observed purchase anchor + public research": 0,
        "Observed purchase adjacency": 1,
        "Public workflow research": 2,
    }
    gaps.sort(key=lambda item: (basis_order[item["evidence_basis"]], item["workflow"], item["missing_category"]))
    return gaps[:30], warnings


def json_write(path: Path, value: Any) -> None:
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(path)


def text_write(path: Path, value: str) -> None:
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(value, encoding="utf-8")
    temporary.replace(path)


def csv_write(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    temporary = path.with_name(f".{path.name}.tmp")
    with temporary.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})
    temporary.replace(path)


def money(value: Any, currency: str) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    prefix = f"{currency} " if currency else ""
    return f"{prefix}{number:,.0f}"


def percent(value: Any) -> str:
    if value in (None, ""):
        return "n.a."
    return f"{float(value) * 100:.1f}%"


def markdown_table(headers: list[str], rows: Iterable[Iterable[Any]]) -> str:
    def safe(value: Any) -> str:
        return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ")
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines.extend("| " + " | ".join(safe(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def research_template(customer: str, catalog: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "customer": customer,
        "company_summary": "",
        "workflows": [],
        "sources": [],
        "available_workflow_ids": [
            {"id": item["id"], "name": item["name"]}
            for item in catalog["workflows"]
        ],
    }


def build_report(summary: dict[str, Any], family_trends: list[dict[str, Any]], gained: list[dict[str, Any]], lost_and_declining: list[dict[str, Any]], lapsed: list[dict[str, Any]], gaps: list[dict[str, Any]], unmapped: list[dict[str, Any]], research: dict[str, Any]) -> str:
    currency = summary["currency"]
    lines = [
        f"# {summary['customer']} GapAnalysis",
        "",
        "## Scope",
        "",
        f"- Source: `{summary['source_file']}` / `{summary['sheet']}`",
        f"- Matched accounts: {'; '.join(summary['matched_accounts'])}",
        f"- Source coverage: {summary['coverage_start']} through {summary['coverage_end']}",
        f"- Comparison: {summary['latest_year']} YTD through {summary['latest_cutoff']} versus {summary['prior_year']} through {summary['prior_cutoff']}",
        f"- Sales measure: {summary['sales_measure_label']}",
        "",
        "## Matched-period results",
        "",
        markdown_table(
            ["Metric", str(summary["latest_year"]), str(summary["prior_year"]), "Change", "Change %"],
            [[
                "Sales",
                money(summary["latest_ytd_sales"], currency),
                money(summary["prior_ytd_sales"], currency),
                money(summary["change"], currency),
                percent(summary["change_percent"]),
            ]],
        ),
        "",
        "## Gained business",
        "",
    ]
    if gained:
        lines.append(markdown_table(
            ["Product family", "Latest YTD", "Prior YTD", "Change"],
            [[row["product_family"], money(row["latest_ytd_sales"], currency), money(row["prior_ytd_sales"], currency), money(row["change"], currency)] for row in gained[:12]],
        ))
    else:
        lines.append("No product family met the mechanical gained-business definition.")
    lines.extend(["", "## Lost and declining business", ""])
    if lost_and_declining:
        lines.append(markdown_table(
            ["Product family", "Status", "Latest YTD", "Prior YTD", "Change"],
            [[row["product_family"], row["movement"], money(row["latest_ytd_sales"], currency), money(row["prior_ytd_sales"], currency), money(row["change"], currency)] for row in lost_and_declining[:15]],
        ))
    else:
        lines.append("No product family showed lost or declining matched-period sales.")
    lines.extend(["", "## Lapsed recurring-product candidates", ""])
    if lapsed:
        lines.append(markdown_table(
            ["SKU", "Product", "Last purchase", "Historical sales", "Why flagged"],
            [[row["sku"], row["product_description"] or row["product_family"], row["last_positive_purchase"], money(row["historical_positive_sales"], currency), f"{row['distinct_purchase_dates']} purchase dates; {row['days_since_last_purchase']} days since last purchase"] for row in lapsed[:15]],
        ))
        lines.append("")
        lines.append("These are win-back hypotheses, not confirmed churn. Confirm project timing, account migration, and workflow continuity.")
    else:
        lines.append("No SKU met the recurring-product lapse rule.")
    lines.extend(["", "## Purchase footprint", ""])
    lines.append(markdown_table(
        ["Product family", "Latest YTD", "Prior YTD", "Change", "Status"],
        [[row["product_family"], money(row["latest_ytd_sales"], currency), money(row["prior_ytd_sales"], currency), money(row["change"], currency), row["movement"]] for row in sorted(family_trends, key=lambda item: -max(abs(item["latest_ytd_sales"]), abs(item["prior_ytd_sales"])))[:20]],
    ))
    lines.extend(["", "## Product gap analysis", ""])
    if gaps:
        lines.append(markdown_table(
            ["Workflow", "Evidence", "Observed anchor", "Category to validate", "Research support"],
            [[row["workflow"], row["evidence_basis"], row["observed_anchor_categories"], row["missing_category"], row["research_evidence"]] for row in gaps[:20]],
        ))
        lines.append("")
        lines.append("Every category above was not observed in the supplied purchase history for the matched account scope and period. This does not prove the customer does not use it or buy it elsewhere.")
    else:
        lines.append("No gap met the purchase-anchor or public-research rule. Add supported workflow research or resolve material unmapped product families before drawing conclusions.")
    if unmapped:
        lines.extend(["", "## Unmapped purchasing", ""])
        lines.append(markdown_table(
            ["Raw product family", "Sales in supplied history"],
            [[row["product_family"], money(row["sales"], currency)] for row in unmapped[:15]],
        ))
        lines.append("")
        lines.append("Review material unmapped families before treating related categories as whitespace.")
    company_summary = str(research.get("company_summary") or "").strip()
    if company_summary:
        lines.extend(["", "## Public research summary", "", company_summary])
    lines.extend([
        "",
        "## Method and limits",
        "",
        "- Latest YTD uses the maximum invoice date in the latest year. Prior YTD uses the exact same inclusive calendar period one year earlier.",
        "- Negative rows and credits remain in sales totals. Blank sales and unparseable dates remain visible in the data-quality table.",
        "- Account sales totals include the selected sales measure as supplied. Freight and other operational-charge rows are excluded from product movement, lapse, and gap tables.",
        "- Gap rows are category-level validation prompts. They are not forecasts, confirmed budgets, or additive pipeline value.",
        "- See the CSV outputs for complete detail and `data-quality.csv` for reconciliation and exceptions.",
        "",
        "## Sources",
        "",
        f"- Internal purchase history: `{summary['source_file']}`, sheet `{summary['sheet']}`.",
    ])
    sources = []
    for item in research.get("workflows", []):
        if item.get("source_url"):
            sources.append(item)
    sources.extend(item for item in research.get("sources", []) if item.get("url") or item.get("source_url"))
    seen = set()
    for item in sources:
        url = str(item.get("url") or item.get("source_url") or "").strip()
        if not url or url in seen:
            continue
        seen.add(url)
        title = str(item.get("title") or item.get("source_title") or "Public source").strip()
        source_date = str(item.get("date") or item.get("source_date") or "").strip()
        label = f"{title} ({source_date})" if source_date else title
        lines.append(f"- [{label}]({url})")
    return "\n".join(lines).rstrip() + "\n"


def run(args: argparse.Namespace) -> dict[str, Any]:
    source = Path(args.input).expanduser().resolve()
    if not source.is_file():
        raise ValueError(f"Input file not found: {source}")
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    script_dir = Path(__file__).resolve().parent
    catalog_path = Path(args.catalog).expanduser().resolve() if args.catalog else script_dir.parent / "references" / "workflow-catalog.json"
    category_map_path = Path(args.category_map).expanduser().resolve() if args.category_map else None
    research_path = Path(args.research).expanduser().resolve() if args.research else None

    profile, raw_rows = read_source(source, args.sheet)
    records, counters, measure_field = build_records(profile, raw_rows)
    identities = account_identities(records)
    resolution = resolve_accounts(args.customer, identities)
    selected_keys = {item["key"] for item in resolution["matches"]}
    selected = [record for record in records if account_key(record) in selected_keys]
    product_records = [record for record in selected if not is_operational(record)]
    operational_records = [record for record in selected if is_operational(record)]
    dated = [record for record in selected if record.get("invoice_date")]
    if not dated:
        raise ValueError("No parseable invoice dates were found for the selected account scope.")
    latest_cutoff = max(record["invoice_date"] for record in dated)
    earliest_date = min(record["invoice_date"] for record in dated)
    previous_cutoff = prior_cutoff(latest_cutoff)
    valid_sales = [record for record in selected if record.get("invoice_date") and record.get("sales") is not None]
    latest_sales = round(sum(float(record["sales"]) for record in valid_sales if period_flags(record["invoice_date"], latest_cutoff)[0]), 2)
    previous_sales = round(sum(float(record["sales"]) for record in valid_sales if period_flags(record["invoice_date"], latest_cutoff)[1]), 2)

    family_trends = aggregate_trends(product_records, latest_cutoff, ("product_family",))
    sku_trends = aggregate_trends(product_records, latest_cutoff, ("sku", "product_description", "product_family"))
    gained = [row for row in family_trends if row["movement"] == "gained"]
    lost = [row for row in family_trends if row["movement"] == "lost"]
    lost_and_declining = [row for row in family_trends if row["movement"] in {"lost", "declining"}]
    lapsed = lapsed_candidates(product_records, latest_cutoff)

    catalog = load_catalog(catalog_path)
    overrides = load_category_map(category_map_path, set(catalog["categories"]))
    current_observed, historical_observed, current_category_sales, historical_category_sales, unmapped_sales = category_evidence(product_records, latest_cutoff, catalog, overrides)
    research = load_research(research_path)
    gaps, research_warnings = build_gaps(catalog, current_observed, historical_observed, research)
    unmapped = [
        {"product_family": family, "sales": round(sales, 2)}
        for family, sales in sorted(unmapped_sales.items(), key=lambda item: (-abs(item[1]), item[0].casefold()))
        if abs(sales) > 0.005
    ]

    currencies = sorted({record["currency"] for record in selected if record.get("currency")})
    currency = currencies[0] if len(currencies) == 1 else ("Multiple" if currencies else "")
    all_selected_sales = round(sum(float(record["sales"]) for record in selected if record.get("sales") is not None), 2)
    undated_sales = round(sum(float(record["sales"]) for record in selected if record.get("sales") is not None and not record.get("invoice_date")), 2)
    operational_sales = round(sum(float(record["sales"]) for record in operational_records if record.get("sales") is not None), 2)
    summary = {
        "schema_version": 1,
        "analyzer_version": VERSION,
        "customer": resolution["display_name"],
        "customer_query": args.customer or "",
        "resolution_method": resolution["method"],
        "matched_accounts": [item["label"] for item in resolution["matches"]],
        "source_file": source.name,
        "sheet": profile.sheet,
        "header_row": profile.header_row,
        "sales_measure": measure_field,
        "sales_measure_label": "Net Sales" if measure_field == "net_sales" else "Extended Sales",
        "currency": currency,
        "coverage_start": earliest_date.isoformat(),
        "coverage_end": latest_cutoff.isoformat(),
        "latest_year": latest_cutoff.year,
        "latest_cutoff": latest_cutoff.isoformat(),
        "prior_year": previous_cutoff.year,
        "prior_cutoff": previous_cutoff.isoformat(),
        "latest_ytd_sales": latest_sales,
        "prior_ytd_sales": previous_sales,
        "change": round(latest_sales - previous_sales, 2),
        "change_percent": pct_change(latest_sales, previous_sales),
        "selected_rows": len(selected),
        "valid_analysis_rows": len(valid_sales),
        "all_selected_sales": all_selected_sales,
        "undated_sales": undated_sales,
        "operational_charge_rows": len(operational_records),
        "operational_charge_sales": operational_sales,
        "gained_family_count": len(gained),
        "lost_family_count": len(lost),
        "declining_family_count": len([row for row in family_trends if row["movement"] == "declining"]),
        "lapsed_candidate_count": len(lapsed),
        "gap_candidate_count": len(gaps),
        "observed_categories_latest_ytd": sorted(current_observed),
        "observed_categories_history": sorted(historical_observed),
        "category_sales_latest_ytd": {key: round(value, 2) for key, value in sorted(current_category_sales.items())},
        "category_sales_history": {key: round(value, 2) for key, value in sorted(historical_category_sales.items())},
    }

    quality = [
        {"check": "account_scope", "status": "PASS", "value": len(resolution["matches"]), "detail": "; ".join(summary["matched_accounts"])},
        {"check": "source_rows", "status": "INFO", "value": counters["source_rows"], "detail": "Rows below the detected header."},
        {"check": "selected_rows", "status": "PASS" if selected else "FAIL", "value": len(selected), "detail": "Rows in the matched account scope."},
        {"check": "missing_sales_selected", "status": "WARN" if any(record.get("sales") is None for record in selected) else "PASS", "value": sum(1 for record in selected if record.get("sales") is None), "detail": "Selected rows excluded from sales totals because the chosen sales measure is blank or invalid."},
        {"check": "missing_date_selected", "status": "WARN" if any(not record.get("invoice_date") for record in selected) else "PASS", "value": sum(1 for record in selected if not record.get("invoice_date")), "detail": f"Undated selected sales: {undated_sales:.2f}."},
        {"check": "sales_reconciliation", "status": "PASS", "value": all_selected_sales, "detail": "Sum of all numeric selected sales rows, including credits and undated rows."},
        {"check": "operational_charges", "status": "INFO", "value": operational_sales, "detail": f"{len(operational_records)} freight, shipping, handling, delivery, or tax rows excluded from product movement, lapse, and gap tables."},
        {"check": "currencies", "status": "WARN" if len(currencies) > 1 else "PASS", "value": "; ".join(currencies), "detail": "Multiple currencies are not converted." if len(currencies) > 1 else ""},
        {"check": "unmapped_product_families", "status": "WARN" if unmapped else "PASS", "value": len(unmapped), "detail": f"Absolute unmapped sales: {sum(abs(item['sales']) for item in unmapped):.2f}."},
        {"check": "public_research", "status": "PASS" if research_path else "INFO", "value": str(research_path or ""), "detail": "Research integrated." if research_path else "No research file supplied; gaps use purchase adjacency only."},
    ]
    quality.extend({"check": "research_workflow", "status": "WARN", "value": "", "detail": warning} for warning in research_warnings)

    family_fields = ["product_family", "latest_ytd_sales", "prior_ytd_sales", "change", "change_percent", "movement", "latest_ytd_quantity", "prior_ytd_quantity", "latest_ytd_rows", "prior_ytd_rows"]
    sku_fields = ["sku", "product_description", "product_family", "latest_ytd_sales", "prior_ytd_sales", "change", "change_percent", "movement", "latest_ytd_quantity", "prior_ytd_quantity", "latest_ytd_rows", "prior_ytd_rows"]
    gap_fields = ["workflow", "evidence_basis", "observed_anchor_categories", "missing_category", "gap_statement", "research_evidence", "source_title", "source_url", "validation_question"]
    lapsed_fields = ["sku", "product_description", "product_family", "distinct_purchase_dates", "last_positive_purchase", "median_cadence_days", "days_since_last_purchase", "historical_positive_sales", "status"]

    csv_write(output_dir / "product-family-trends.csv", family_trends, family_fields)
    csv_write(output_dir / "sku-trends.csv", sku_trends, sku_fields)
    csv_write(output_dir / "gained-business.csv", gained, family_fields)
    csv_write(output_dir / "lost-business.csv", lost_and_declining, family_fields)
    csv_write(output_dir / "lapsed-business.csv", lapsed, lapsed_fields)
    csv_write(output_dir / "gap-analysis.csv", gaps, gap_fields)
    csv_write(output_dir / "unmapped-product-families.csv", unmapped, ["product_family", "sales"])
    csv_write(output_dir / "data-quality.csv", quality, ["check", "status", "value", "detail"])
    json_write(output_dir / "gapanalysis-summary.json", summary)
    text_write(output_dir / "gapanalysis-report.md", build_report(summary, family_trends, gained, lost_and_declining, lapsed, gaps, unmapped, research))
    if not research_path:
        json_write(output_dir / "research-template.json", research_template(summary["customer"], catalog))

    return {
        "status": "ok",
        "customer": summary["customer"],
        "matched_accounts": summary["matched_accounts"],
        "output_dir": str(output_dir),
        "report": str(output_dir / "gapanalysis-report.md"),
        "latest_ytd_sales": latest_sales,
        "prior_ytd_sales": previous_sales,
        "gained_families": len(gained),
        "lost_or_declining_families": len(lost_and_declining),
        "lapsed_candidates": len(lapsed),
        "gap_candidates": len(gaps),
        "warnings": [item["detail"] for item in quality if item["status"] == "WARN"],
    }


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description="Analyze customer purchase history and product-category gaps without charts or commercial scoring.")
    result.add_argument("input", help="Source .xlsx, .xlsm, or .csv purchase-history file")
    result.add_argument("--customer", help="Customer name, exact account label, or exact numeric SGN/account id")
    result.add_argument("--output-dir", required=True, help="Directory for the generated report and CSV tables")
    result.add_argument("--sheet", help="Worksheet name when auto-detection is ambiguous")
    result.add_argument("--research", help="Optional research JSON created from research-template.json")
    result.add_argument("--category-map", help="Optional CSV with match,category columns for local product-family mapping")
    result.add_argument("--catalog", help="Optional replacement workflow catalog JSON")
    return result


def main() -> int:
    args = parser().parse_args()
    try:
        outcome = run(args)
    except (ValueError, OSError, json.JSONDecodeError, csv.Error) as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, indent=2), file=sys.stderr)
        return 2
    print(json.dumps(outcome, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
