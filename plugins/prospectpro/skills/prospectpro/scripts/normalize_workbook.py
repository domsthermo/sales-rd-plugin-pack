#!/usr/bin/env python3
"""Profile and normalize purchase-history workbooks for ProspectPro.

The script performs mechanical normalization only. It never invents product
families, merges unrequested accounts, or edits the source workbook.
"""

from __future__ import annotations

import argparse
import calendar
import hashlib
import json
import math
import os
import re
import statistics
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


NORMALIZER_VERSION = "5.0"  # Bump whenever parsing or output semantics change.

OPERATIONAL_CHARGE_PHRASES = (
    "freight",
    "shipping charge",
    "hazardous material charge",
    "delivery charge",
    "tax charge",
)


FIELD_ALIASES = {
    "invoice_date": {"invoice date", "invoice dt", "billing date"},
    "order_date": {"order date", "sales order date"},
    "account_name": {"account", "account name", "customer", "customer name", "ship to", "ship to name"},
    "account_id": {"account id", "account number", "account no", "customer id", "customer number", "nsgn"},
    "site_id": {"site id", "site number", "site sgn", "sgn"},
    "sku": {"sku", "catalog number", "catalog no", "catalog #", "product number", "item number", "material number"},
    "product_description": {"product description", "item description", "sku description", "material description", "description"},
    "product_family": {"product family", "product line", "product line group", "product group", "pricing product line", "dsr product group na"},
    "quantity": {"quantity", "qty", "ship quantity", "invoice quantity", "ordered quantity"},
    "unit_price": {"unit price", "net unit price", "price each"},
    "net_sales": {"net sales", "net revenue", "invoice net", "sales net"},
    "extended_sales": {"extended sales", "ext sales", "extended amount", "sales amount", "revenue"},
    "freight": {"freight", "shipping", "shipping charge"},
    "currency": {"currency", "currency code", "curr"},
    "invoice_number": {"sales invoice number", "sales invoice no", "sales invoice #", "invoice number", "invoice #"},
}

CORE_FIELDS = {"invoice_date", "net_sales", "extended_sales", "sku", "account_name", "account_id", "site_id"}

STANDARD_TERRITORY_SCHEMA_NAME = "territory_invoice_report_v1"
STANDARD_TERRITORY_REQUIRED = {
    "account_name": "Ship To",
    "account_id": "NSGN",
    "site_id": "SGN",
    "invoice_date": "Invoice Date",
    "sku": "SKU",
    "product_family": "Pricing Product Line",
    "quantity": "Ship Quantity",
    "net_sales": "Net Sales",
}
STANDARD_TERRITORY_DIAGNOSTIC = {"invoice_number": "Sales Invoice #"}
STANDARD_TERRITORY_SIGNATURE_HEADERS = {
    "NSGN",
    "SGN",
    "Product Line Group",
    "DSR Product Group - NA",
    "Pricing Product Line",
    "Comp Component",
    "BD-RR",
}


def normalized_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def is_operational_charge(*values: Any) -> bool:
    text = " ".join(normalized_text(value) for value in values)
    return any(phrase in text for phrase in OPERATIONAL_CHARGE_PHRASES)


NORMALIZED_ALIASES = {
    field: {normalized_text(alias) for alias in aliases}
    for field, aliases in FIELD_ALIASES.items()
}


def load_overrides(path: str | None) -> dict[str, Any]:
    if not path:
        return {}
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Column map must be a JSON object of field-to-header mappings.")
    unknown = sorted(set(data) - set(FIELD_ALIASES))
    if unknown:
        raise ValueError(f"Unknown canonical fields in column map: {', '.join(unknown)}")
    return data


def file_sha256(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def request_fingerprint(workbook_sha256: str, request: dict[str, Any]) -> str:
    payload = {
        "normalizer_version": NORMALIZER_VERSION,
        "workbook_sha256": workbook_sha256,
        "request": request,
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def write_json_atomic(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    temporary.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(path)


def resolve_headers(headers: list[Any], overrides: dict[str, Any]) -> tuple[dict[str, int], dict[str, list[str]]]:
    normalized_headers = [normalized_text(value) for value in headers]
    mapping: dict[str, int] = {}
    ambiguous: dict[str, list[str]] = {}

    for field, override in overrides.items():
        if isinstance(override, int):
            index = override - 1
            if not 0 <= index < len(headers):
                raise ValueError(f"Column override for {field} is outside the worksheet: {override}")
            mapping[field] = index
            continue
        target = normalized_text(override)
        matches = [index for index, header in enumerate(normalized_headers) if header == target]
        if len(matches) != 1:
            raise ValueError(f"Column override for {field} matched {len(matches)} headers: {override!r}")
        mapping[field] = matches[0]

    for field, aliases in NORMALIZED_ALIASES.items():
        if field in mapping:
            continue
        matches = [index for index, header in enumerate(normalized_headers) if header in aliases]
        if len(matches) == 1:
            mapping[field] = matches[0]
        elif len(matches) > 1:
            ambiguous[field] = [str(headers[index]) for index in matches]

    return mapping, ambiguous


def resolve_standard_territory_schema(
    headers: list[Any],
    schema_mode: str,
) -> tuple[dict[str, int] | None, str | None]:
    if schema_mode == "generic":
        return None, None

    positions: dict[str, list[int]] = defaultdict(list)
    for index, header in enumerate(headers):
        if header not in (None, ""):
            positions[str(header).strip()].append(index)

    signature_matches = len(STANDARD_TERRITORY_SIGNATURE_HEADERS.intersection(positions))
    standard_like = signature_matches >= 3
    if schema_mode == "territory-invoice":
        standard_like = True
    if not standard_like:
        return None, None

    missing = [header for header in STANDARD_TERRITORY_REQUIRED.values() if header not in positions]
    duplicates = [header for header in STANDARD_TERRITORY_REQUIRED.values() if len(positions.get(header, [])) > 1]
    if missing or duplicates:
        details = []
        if missing:
            details.append(f"missing required columns: {', '.join(missing)}")
        if duplicates:
            details.append(f"duplicate required columns: {', '.join(duplicates)}")
        raise ValueError(
            "The workbook resembles the standard Territory Invoice Report but does not match its exact schema; "
            + "; ".join(details)
        )

    mapping = {field: positions[header][0] for field, header in STANDARD_TERRITORY_REQUIRED.items()}
    for field, header in STANDARD_TERRITORY_DIAGNOSTIC.items():
        if len(positions.get(header, [])) == 1:
            mapping[field] = positions[header][0]
    return mapping, STANDARD_TERRITORY_SCHEMA_NAME


def header_score(mapping: dict[str, int]) -> int:
    score = len(mapping) * 10
    score += sum(8 for field in CORE_FIELDS if field in mapping)
    if "invoice_date" in mapping:
        score += 20
    if "net_sales" in mapping or "extended_sales" in mapping:
        score += 20
    return score


def profile_sheets(workbook: Any, overrides: dict[str, Any], scan_rows: int = 30) -> list[dict[str, Any]]:
    profiles = []
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        candidates = []
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(scan_rows, worksheet.max_row), values_only=True),
            start=1,
        ):
            headers = list(row)
            mapping, ambiguous = resolve_headers(headers, overrides)
            candidates.append(
                {
                    "row": row_number,
                    "score": header_score(mapping),
                    "headers": headers,
                    "mapping": mapping,
                    "ambiguous": ambiguous,
                }
            )
        best = max(candidates, key=lambda item: (item["score"], -item["row"]))
        profiles.append(
            {
                "sheet": sheet_name,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "header_row": best["row"],
                "header_score": best["score"],
                "headers": best["headers"],
                "mapping": best["mapping"],
                "ambiguous": best["ambiguous"],
            }
        )
    return profiles


def choose_sheet(profiles: list[dict[str, Any]], requested: str | None) -> dict[str, Any]:
    if requested:
        matches = [profile for profile in profiles if profile["sheet"] == requested]
        if not matches:
            raise ValueError(f"Worksheet not found: {requested}")
        return matches[0]
    best_score = max(profile["header_score"] for profile in profiles)
    matches = [profile for profile in profiles if profile["header_score"] == best_score]
    if best_score < 40:
        raise ValueError("No worksheet contains a credible purchase-history header row.")
    if len(matches) > 1:
        names = ", ".join(profile["sheet"] for profile in matches)
        raise ValueError(f"Multiple worksheets are equally plausible. Specify --sheet: {names}")
    return matches[0]


def value_at(row: tuple[Any, ...], mapping: dict[str, int], field: str) -> Any:
    index = mapping.get(field)
    return row[index] if index is not None and index < len(row) else None


def parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"[$,£€%()]", "", text).strip()
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


DATE_FORMATS = (
    "%Y-%m-%d",
    "%m/%d/%Y",
    "%m/%d/%y",
    "%Y/%m/%d",
    "%d-%b-%Y",
    "%d-%b-%y",
)


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        integer_value = int(value)
        if float(value).is_integer() and 10_000_101 <= integer_value <= 99_991_231:
            try:
                return datetime.strptime(str(integer_value), "%Y%m%d").date()
            except ValueError:
                return None
        try:
            parsed = from_excel(value)
            return parsed.date() if isinstance(parsed, datetime) else parsed
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


def prior_year_cutoff(cutoff: date) -> date:
    prior_year = cutoff.year - 1
    day = min(cutoff.day, calendar.monthrange(prior_year, cutoff.month)[1])
    return date(prior_year, cutoff.month, day)


def clean_identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def analyze(
    workbook_path: Path,
    sheet: str | None,
    accounts: list[str] | None,
    account_field_requested: str | None,
    schema_mode: str,
    overrides: dict[str, Any],
) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    profiles = profile_sheets(workbook, {})
    selected = choose_sheet(profiles, sheet)
    worksheet = workbook[selected["sheet"]]
    headers = selected["headers"]
    standard_mapping, selected_schema = resolve_standard_territory_schema(headers, schema_mode)
    if standard_mapping is not None:
        if overrides:
            raise ValueError("Column-map overrides are not allowed for the standard Territory Invoice Report schema.")
        if account_field_requested not in (None, "site_id"):
            raise ValueError("The standard Territory Invoice Report always uses SGN as the customer filter.")
        mapping = standard_mapping
        ambiguous: dict[str, list[str]] = {}
    else:
        mapping, ambiguous = resolve_headers(headers, overrides)
    selected["mapping"] = mapping
    selected["ambiguous"] = ambiguous

    sales_field = "net_sales" if "net_sales" in mapping else "extended_sales" if "extended_sales" in mapping else None
    if "invoice_date" not in mapping or not sales_field:
        raise ValueError("The selected sheet needs an invoice-date column and either net-sales or extended-sales column.")
    if "sku" not in mapping and "product_description" not in mapping:
        raise ValueError("The selected sheet needs a SKU/catalog-number column or a product-description column.")

    warnings: list[str] = []
    if sales_field != "net_sales":
        warnings.append("Net Sales was not found; Extended Sales is used and must be labeled accordingly.")
    for field, candidates in selected["ambiguous"].items():
        warnings.append(f"Ambiguous {field} columns were not selected automatically: {', '.join(candidates)}")

    if account_field_requested and account_field_requested not in mapping:
        raise ValueError(f"Requested account field is not mapped: {account_field_requested}")
    if selected_schema == STANDARD_TERRITORY_SCHEMA_NAME:
        account_field = "site_id"
    else:
        account_field = account_field_requested or next(
            (field for field in ("account_name", "account_id", "site_id") if field in mapping),
            None,
        )
    account_targets = {
        normalized_text(account)
        for account in (accounts or [])
        if normalized_text(account)
    }
    account_counts: Counter[str] = Counter()
    account_labels: dict[str, str] = {}
    rows: list[dict[str, Any]] = []
    raw_sales = 0.0
    raw_fact_rows = 0
    missing_dates = 0
    missing_sales = 0

    for row in worksheet.iter_rows(min_row=selected["header_row"] + 1, values_only=True):
        if not any(value not in (None, "") for value in row):
            continue

        account_value = clean_identifier(value_at(row, mapping, account_field)) if account_field else ""
        account_key = normalized_text(account_value)
        if account_key:
            account_counts[account_key] += 1
            account_labels.setdefault(account_key, account_value)
        if account_targets and account_key not in account_targets:
            continue

        raw_fact_rows += 1
        sales = parse_number(value_at(row, mapping, sales_field))
        invoice_date = parse_date(value_at(row, mapping, "invoice_date"))
        if sales is None:
            missing_sales += 1
        else:
            raw_sales += sales
        if invoice_date is None:
            missing_dates += 1
        if sales is None or invoice_date is None:
            continue

        sku = clean_identifier(value_at(row, mapping, "sku"))
        description = clean_identifier(value_at(row, mapping, "product_description"))
        family = clean_identifier(value_at(row, mapping, "product_family"))
        quantity = parse_number(value_at(row, mapping, "quantity"))
        group_label = family or sku or description or "Unmapped"
        rows.append(
            {
                "date": invoice_date,
                "sales": sales,
                "quantity": quantity,
                "sku": sku,
                "description": description,
                "family": family,
                "group_label": group_label,
            }
        )

    known_accounts = [
        {"value": account_labels[key], "row_count": count}
        for key, count in account_counts.most_common()
    ]
    if not account_targets and len(known_accounts) > 1:
        warnings.append(
            "Multiple account values are present. Use this result only to select one or more exact account identifiers, then rerun with --account."
        )
    if "quantity" not in mapping:
        warnings.append("No quantity column was mapped; sales analysis can proceed, but unit and cadence analysis is limited.")
    if "product_family" not in mapping and "product_family" not in selected["ambiguous"]:
        warnings.append("No product-family column was mapped; product rows are grouped by SKU or description.")
    if account_targets and raw_fact_rows == 0:
        raise ValueError(
            "No rows matched the supplied exact account values. Use the reported workbook account values exactly."
        )
    if not rows:
        raise ValueError("No usable rows remain after account, date, and sales validation.")

    latest_date = max(item["date"] for item in rows)
    latest_year = latest_date.year
    previous_cutoff = prior_year_cutoff(latest_date)
    latest_start = date(latest_year, 1, 1)
    previous_start = date(latest_year - 1, 1, 1)

    grouped: dict[tuple[str, str, str], dict[str, Any]] = defaultdict(
        lambda: {
            "latest_ytd_sales": 0.0,
            "prior_ytd_sales": 0.0,
            "latest_ytd_quantity": 0.0,
            "prior_ytd_quantity": 0.0,
            "positive_purchase_dates": set(),
            "positive_sales": 0.0,
            "credit_sales": 0.0,
        }
    )
    latest_total = 0.0
    prior_total = 0.0
    included_sales = 0.0

    for item in rows:
        included_sales += item["sales"]
        key = (item["group_label"], item["sku"], item["description"])
        if item["sales"] > 0:
            grouped[key]["positive_purchase_dates"].add(item["date"])
            grouped[key]["positive_sales"] += item["sales"]
        elif item["sales"] < 0:
            grouped[key]["credit_sales"] += item["sales"]
        if latest_start <= item["date"] <= latest_date:
            latest_total += item["sales"]
            grouped[key]["latest_ytd_sales"] += item["sales"]
            if item["quantity"] is not None:
                grouped[key]["latest_ytd_quantity"] += item["quantity"]
        elif previous_start <= item["date"] <= previous_cutoff:
            prior_total += item["sales"]
            grouped[key]["prior_ytd_sales"] += item["sales"]
            if item["quantity"] is not None:
                grouped[key]["prior_ytd_quantity"] += item["quantity"]

    product_rows = []
    lapsed_products = []
    lapsed_operational_excluded = 0
    for (group_label, sku, description), values in grouped.items():
        if not values["latest_ytd_sales"] and not values["prior_ytd_sales"]:
            continue
        product_rows.append(
            {
                "group": group_label,
                "sku": sku or None,
                "description": description or None,
                "latest_ytd_sales": round(values["latest_ytd_sales"], 2),
                "prior_ytd_sales": round(values["prior_ytd_sales"], 2),
                "latest_ytd_quantity": round(values["latest_ytd_quantity"], 2),
                "prior_ytd_quantity": round(values["prior_ytd_quantity"], 2),
            }
        )

        purchase_dates = sorted(values["positive_purchase_dates"])
        if account_targets and len(purchase_dates) >= 3 and values["latest_ytd_sales"] <= 0:
            if is_operational_charge(group_label, sku, description):
                lapsed_operational_excluded += 1
                continue
            intervals = [
                (later - earlier).days
                for earlier, later in zip(purchase_dates, purchase_dates[1:])
                if later > earlier
            ]
            if len(intervals) >= 2:
                median_interval = float(statistics.median(intervals))
                lapse_threshold = max(90, math.ceil(median_interval * 1.75))
                days_since_last = (latest_date - purchase_dates[-1]).days
                if days_since_last > lapse_threshold:
                    lapsed_products.append(
                        {
                            "group": group_label,
                            "sku": sku or None,
                            "description": description or None,
                            "classification": "lapsed_candidate",
                            "positive_purchase_date_count": len(purchase_dates),
                            "first_positive_purchase_date": purchase_dates[0].isoformat(),
                            "last_positive_purchase_date": purchase_dates[-1].isoformat(),
                            "median_reorder_interval_days": round(median_interval, 1),
                            "lapse_threshold_days": lapse_threshold,
                            "days_since_last_positive_purchase": days_since_last,
                            "historical_positive_sales": round(values["positive_sales"], 2),
                            "historical_credits": round(values["credit_sales"], 2),
                            "latest_ytd_sales": round(values["latest_ytd_sales"], 2),
                            "prior_ytd_sales": round(values["prior_ytd_sales"], 2),
                            "requires_product_type_review": True,
                        }
                    )
    product_rows.sort(
        key=lambda item: (abs(item["latest_ytd_sales"]) + abs(item["prior_ytd_sales"]), item["group"]),
        reverse=True,
    )
    lapsed_products.sort(
        key=lambda item: (item["prior_ytd_sales"], item["historical_positive_sales"], item["group"]),
        reverse=True,
    )

    controlling_fields = [
        field
        for field in (account_field, "invoice_date", sales_field, "sku")
        if field and field in mapping
    ]
    supporting_fields = [
        field
        for field in ("product_description", "product_family", "quantity", "currency")
        if field in mapping
    ]
    identity_check_fields = [
        field
        for field in ("account_name", "account_id", "site_id")
        if field in mapping and field != account_field
    ]
    diagnostic_fields = [field for field in ("invoice_number",) if field in mapping]
    used_fields = set(controlling_fields + supporting_fields + identity_check_fields + diagnostic_fields)
    used_indices = {mapping[field] for field in used_fields}
    unused_headers = [
        str(header)
        for index, header in enumerate(headers)
        if header not in (None, "") and index not in used_indices
    ]

    return {
        "source": {
            "file": str(workbook_path.resolve()),
            "sheet": selected["sheet"],
            "header_row": selected["header_row"],
            "account_filter": accounts or None,
            "account_field": account_field,
            "schema": selected_schema or "generic",
            "sales_measure": sales_field,
        },
        "sheet_profiles": [
            {
                "sheet": profile["sheet"],
                "max_row": profile["max_row"],
                "max_column": profile["max_column"],
                "header_row": profile["header_row"],
                "header_score": profile["header_score"],
            }
            for profile in profiles
        ],
        "column_mapping": {field: str(headers[index]) for field, index in mapping.items()},
        "column_usage": {
            "controlling": {field: str(headers[mapping[field]]) for field in controlling_fields},
            "supporting": {field: str(headers[mapping[field]]) for field in supporting_fields},
            "identity_checks": {field: str(headers[mapping[field]]) for field in identity_check_fields},
            "diagnostic": {field: str(headers[mapping[field]]) for field in diagnostic_fields},
            "available_but_not_used": unused_headers,
        },
        "ambiguous_columns": selected["ambiguous"],
        "known_accounts": known_accounts,
        "coverage": {
            "usable_row_count": len(rows),
            "raw_fact_row_count": raw_fact_rows,
            "earliest_invoice_date": min(item["date"] for item in rows).isoformat(),
            "latest_invoice_date": latest_date.isoformat(),
            "latest_ytd_start": latest_start.isoformat(),
            "latest_ytd_cutoff": latest_date.isoformat(),
            "prior_ytd_start": previous_start.isoformat(),
            "prior_ytd_cutoff": previous_cutoff.isoformat(),
        },
        "reconciliation": {
            "raw_sales": round(raw_sales, 2),
            "included_sales_with_valid_date": round(included_sales, 2),
            "difference": round(raw_sales - included_sales, 2),
            "rows_missing_date": missing_dates,
            "rows_missing_sales": missing_sales,
        },
        "ytd": {
            "latest_year": latest_year,
            "latest_ytd_sales": round(latest_total, 2),
            "prior_year": latest_year - 1,
            "prior_ytd_sales": round(prior_total, 2),
        },
        "product_rows": product_rows,
        "lapsed_detection": {
            "evaluated": bool(account_targets),
            "candidate_count": len(lapsed_products),
            "operational_charge_groups_excluded": lapsed_operational_excluded,
            "minimum_positive_purchase_dates": 3,
            "minimum_lapse_threshold_days": 90,
            "cadence_multiplier": 1.75,
            "requires_zero_or_negative_latest_ytd_sales": True,
            "note": (
                "Candidates require product-type review to exclude instruments, service, project-only, and seasonal purchases."
                if account_targets
                else "Lapsed detection is deferred until one or more exact account identifiers are supplied."
            ),
        },
        "lapsed_products": lapsed_products,
        "warnings": warnings,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", help="Path to the source .xlsx workbook")
    parser.add_argument("--output", required=True, help="Path for the normalized JSON result")
    parser.add_argument("--sheet", help="Exact worksheet name; auto-detected when omitted")
    parser.add_argument(
        "--account",
        action="append",
        help="Exact account value from the selected account column; repeat to aggregate multiple identifiers",
    )
    parser.add_argument(
        "--schema",
        choices=("auto", "territory-invoice", "generic"),
        default="auto",
        help="Schema behavior; auto applies the strict standard report schema when its signature is detected",
    )
    parser.add_argument(
        "--account-field",
        choices=("account_name", "account_id", "site_id"),
        help="Canonical account field used by --account; the standard Territory Invoice Report always uses site_id (SGN)",
    )
    parser.add_argument("--column-map", help="Optional JSON file mapping canonical fields to exact headers or 1-based columns")
    parser.add_argument(
        "--cache-dir",
        help="Private backend directory for fingerprint-keyed normalized results; omit to disable cache reuse",
    )
    args = parser.parse_args()

    try:
        workbook_path = Path(args.workbook).expanduser().resolve()
        if not workbook_path.is_file():
            raise ValueError(f"Workbook not found: {workbook_path}")
        overrides = load_overrides(args.column_map)
        workbook_digest = file_sha256(workbook_path)
        request = {
            "sheet": args.sheet,
            "accounts": args.account,
            "account_field": args.account_field,
            "schema": args.schema,
            "column_map": overrides,
        }
        fingerprint = request_fingerprint(workbook_digest, request)
        cache_path = Path(args.cache_dir).expanduser().resolve() / f"{fingerprint}.json" if args.cache_dir else None
        cache_hit = False
        result = None
        if cache_path and cache_path.is_file():
            try:
                candidate = json.loads(cache_path.read_text(encoding="utf-8"))
                if (
                    candidate.get("backend", {}).get("normalization_fingerprint") == fingerprint
                    and isinstance(candidate.get("coverage"), dict)
                    and isinstance(candidate.get("source"), dict)
                ):
                    result = candidate
                    result["source"]["file"] = str(workbook_path)
                    cache_hit = True
            except (OSError, json.JSONDecodeError, AttributeError):
                result = None
        if result is None:
            result = analyze(
                workbook_path,
                args.sheet,
                args.account,
                args.account_field,
                args.schema,
                overrides,
            )

        result["backend"] = {
            "normalizer_version": NORMALIZER_VERSION,
            "workbook_sha256": workbook_digest,
            "normalization_fingerprint": fingerprint,
            "cache_enabled": cache_path is not None,
            "cache_hit": cache_hit,
            "workbook_size_bytes": workbook_path.stat().st_size,
        }
        if cache_path and not cache_hit:
            cached_result = json.loads(json.dumps(result))
            cached_result["backend"]["cache_hit"] = False
            write_json_atomic(cache_path, cached_result)

        output_path = Path(args.output)
        write_json_atomic(output_path, result)
    except Exception as exc:
        print(f"ProspectPro workbook normalization failed: {exc}", file=sys.stderr)
        return 2

    action = "Reused" if result["backend"]["cache_hit"] else "Normalized"
    print(
        f"{action} {result['coverage']['usable_row_count']} rows from {result['source']['sheet']} "
        f"through {result['coverage']['latest_ytd_cutoff']} "
        f"(fingerprint {result['backend']['normalization_fingerprint'][:12]})."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
